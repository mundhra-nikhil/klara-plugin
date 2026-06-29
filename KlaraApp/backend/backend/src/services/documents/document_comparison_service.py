"""Before/after document comparison + processed-document generation.

The AI engine emits per-rule findings carrying ``original_text`` and
``replacement_text`` (plus an operator-edited ``applied_text`` once a QC
operator accepts/edits a finding). This service turns those corrections into:

  * a side-by-side / redline HTML comparison for the in-app viewer, and
  * a real corrected ``.docx`` for download.

By default it reflects only the corrections a QC operator has **accepted**
(``applied_text`` overrides the AI's suggestion). If QC hasn't reviewed yet, it
falls back to every AI-suggested replacement so the demo still shows an
"after" document.
"""

import difflib
import html as html_mod
import os
import re
import shutil
from io import BytesIO
from uuid import UUID

from sqlalchemy.ext.asyncio import AsyncSession

from src.models.enum.finding_status import FindingStatus
from src.repositories import qc_finding_repo
from src.services.ai_jobs.poc_rules import TRACK_LABELS, RULE_NAME, poc_result
from src.utils.text_extractor import (
    docx_to_html,
    docx_to_resolved_paragraphs,
    estimate_page_map,
)
from src.core.logger import get_logger_with_context

logger = get_logger_with_context()

WORD_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _project_root() -> str:
    # .../epiq/backend/src/services/documents/<this file>  →  .../epiq
    here = os.path.dirname(os.path.abspath(__file__))
    return os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(here))))


def _backend_root() -> str:
    here = os.path.dirname(os.path.abspath(__file__))
    return os.path.dirname(os.path.dirname(os.path.dirname(here)))


def resolve_document_file(doc) -> str:
    """Best-effort resolution of the saved document file on disk.

    Mirrors the path-resolution used by /documents/{id}/content and the AI job
    worker: documents uploaded via /upload-local set blob_url to
    "/uploads/<filename>"; otherwise we match by title in known directories.
    Returns "" if no file can be located."""
    project_root = _project_root()
    backend_root = _backend_root()

    candidates: list[str] = []
    blob_url = getattr(doc, "blob_url", None)
    if isinstance(blob_url, str) and blob_url.startswith("/uploads/"):
        base = os.path.basename(blob_url)
        candidates.append(os.path.join("/uploads", base))
        candidates.append(os.path.join(project_root, "frontend", "uploads", base))

    title = getattr(doc, "title", None)
    if title:
        for d in [
            "/uploads",
            "/test_files",
            os.path.join(project_root, "frontend", "uploads"),
            os.path.join(project_root, "test_files"),
            os.path.join(project_root, "documentation"),
            backend_root,
            project_root,
        ]:
            candidates.append(os.path.join(d, title))

    return next((p for p in candidates if os.path.exists(p)), "")


# ── Pristine-original preservation ────────────────────────────────────────────
# The OnlyOffice save-callback overwrites the working file on disk every time the
# reviewer saves, so the only way to diff "what changed" is to keep an untouched
# snapshot of the document as it was *before* editing began. We stash that copy
# in an ``originals/`` folder beside the working file and never overwrite it once
# created. ``resolve_original_file`` falls back to the working file for documents
# that predate this feature (their diff simply comes out empty).

def _original_backup_path(working_path: str) -> str:
    """Path of the pristine snapshot for a given working file."""
    d, base = os.path.split(working_path)
    return os.path.join(d, "originals", base)


def ensure_original_backup(working_path: str, refresh: bool = False) -> str | None:
    """Make sure a pristine snapshot of ``working_path`` exists.

    ``refresh=True`` (fresh upload) rewrites the snapshot to match the new file —
    this resets the baseline. ``refresh=False`` (about to overwrite via an editor
    save) only creates the snapshot if one isn't there yet, so the very first
    pre-edit state is preserved and never clobbered. Returns the snapshot path, or
    None on failure."""
    if not working_path or not os.path.exists(working_path):
        return None
    backup = _original_backup_path(working_path)
    try:
        if refresh or not os.path.exists(backup):
            os.makedirs(os.path.dirname(backup), exist_ok=True)
            shutil.copy2(working_path, backup)
            logger.info("Saved pristine original snapshot", backup=backup, refresh=refresh)
        return backup
    except Exception as e:
        logger.error(f"Failed to snapshot original for {working_path}: {e}")
        return None


def resolve_original_file(doc) -> str:
    """Best-effort path to the pristine original. Falls back to the current
    working file when no snapshot exists (documents uploaded before snapshots
    were introduced) so callers always get *something* to diff against."""
    working = resolve_document_file(doc)
    if not working:
        return ""
    backup = _original_backup_path(working)
    return backup if os.path.exists(backup) else working


def _derive_replacement(rule_id, original: str) -> str | None:
    """Deterministically derive a corrected string for the text-content rules
    the engine can fix verbatim, when the LLM left ``replacement_text`` empty.

    Rule 5 — two spaces after a period. Rule 8 — one space after § symbol.
    Returns None for rules whose fix isn't a verbatim text substitution
    (e.g. font size / line spacing properties)."""
    if not original:
        return None
    if rule_id == 5:
        # Insert two spaces after a period that's directly followed by a letter.
        fixed = re.sub(r"\.(?=[A-Za-z])", ".  ", original)
        return fixed if fixed != original else None
    if rule_id == 8:
        # Ensure exactly one space after the section symbol.
        fixed = re.sub(r"§(?=\S)", "§ ", original)
        return fixed if fixed != original else None
    return None


# Deterministic, document-wide text transforms for rules whose fix is an
# unambiguous global substitution. Applied across the whole document whenever a
# finding for that rule exists — the LLM's quoted ``original_text`` is frequently
# paraphrased or concatenates non-adjacent references (e.g. "§17200, §1709"),
# so a verbatim find-and-replace silently drops it. Applying the rule's own
# transform to the document text instead makes the correction actually land.
# Only rules whose transform is safe to apply blindly belong here (NOT rule 5 —
# "two spaces after a period" would wrongly expand abbreviations like "U.S.").
_RULE_REGEX_TRANSFORMS: dict[int, tuple[re.Pattern, str]] = {
    8: (re.compile(r"§(?=\S)"), "§ "),  # exactly one space after the section symbol
}


def _rule_transforms_for(corrections: list[dict]) -> list[tuple[re.Pattern, str]]:
    """The deterministic document-wide transforms implied by the rules present in
    ``corrections`` (deduped, in rule order)."""
    rules = sorted({c.get("rule_id") for c in corrections if c.get("rule_id") in _RULE_REGEX_TRANSFORMS})
    return [_RULE_REGEX_TRANSFORMS[r] for r in rules]


def _correction_from_finding(finding) -> dict | None:
    """Build a {original, replacement, ...} correction from a QCFinding, or
    None if it doesn't carry an actionable text replacement."""
    loc = finding.location if isinstance(finding.location, dict) else {}
    original = loc.get("original_text")
    rule_id = loc.get("rule_id")
    # Operator's inline edit wins over the AI's suggested replacement.
    replacement = loc.get("applied_text")
    if replacement is None:
        replacement = loc.get("replacement_text")
    # The LLM sometimes omits replacement_text (empty string) even though it
    # flagged a verbatim text issue — derive a deterministic fix for those.
    if (replacement is None or replacement == "") and original:
        replacement = _derive_replacement(rule_id, original)
    # An empty replacement must NOT be applied (it would delete the original).
    if not original or not replacement or original == replacement:
        return None

    track = loc.get("track")
    return {
        "rule_id": rule_id,
        "rule_name": loc.get("rule_name") or RULE_NAME.get(rule_id),
        "track": track,
        "track_label": TRACK_LABELS.get(track) if track else None,
        "original": original,
        "replacement": replacement,
        "paragraph": loc.get("paragraph"),
        # Strict POC classification (pass/fail/warning) — no severity.
        "result": loc.get("result") or poc_result(rule_id),
        "type": getattr(finding.finding_type, "value", None),
        "description": finding.description,
    }


async def collect_corrections(db: AsyncSession, document_id: UUID) -> list[dict]:
    """Gather the corrections to apply to the document. QC-accepted findings
    take precedence; if none are accepted yet, fall back to all AI-suggested
    replacements."""
    findings = await qc_finding_repo.find_by_document_id(db, document_id)
    accepted = [f for f in findings if f.status == FindingStatus.ACCEPTED]
    source = accepted if accepted else findings

    corrections: list[dict] = []
    seen: set[tuple] = set()
    for f in source:
        c = _correction_from_finding(f)
        if not c:
            continue
        key = (c["original"], c["replacement"])
        if key in seen:
            continue
        seen.add(key)
        corrections.append(c)
    return corrections


async def collect_rule_findings(
    db: AsyncSession, document_id: UUID, applied_keys: set[tuple]
) -> list[dict]:
    """Build a detailed per-rule view of EVERY AI finding (not just the text
    substitutions applied to the processed doc). Drives the "Compare original
    vs AI" detailed breakdown: each entry carries the rule, its POC result
    (fail/warning), the offending text, the proposed correction, the location,
    and whether that correction was applied to the processed document."""
    findings = await qc_finding_repo.find_by_document_id(db, document_id)
    out: list[dict] = []
    for f in findings:
        loc = f.location if isinstance(f.location, dict) else {}
        rid = loc.get("rule_id")
        original = loc.get("original_text")
        replacement = loc.get("applied_text")
        if replacement is None:
            replacement = loc.get("replacement_text")
        if (replacement is None or replacement == "") and original:
            replacement = _derive_replacement(rid, original)
        track = loc.get("track")
        applied = bool(original and replacement and (original, replacement) in applied_keys)
        out.append({
            "rule_id": rid,
            "rule_name": loc.get("rule_name") or RULE_NAME.get(rid) or "—",
            "track": track,
            "track_label": TRACK_LABELS.get(track) if track else None,
            "result": loc.get("result") or poc_result(rid),
            "original": original,
            "replacement": replacement,
            "suggested_fix": f.suggested_fix,
            "paragraph": loc.get("paragraph"),
            "description": f.description,
            "status": getattr(f.status, "value", None),
            "applied": applied,
        })
    out.sort(key=lambda r: (r["rule_id"] is None, r["rule_id"] or 0))
    return out


def result_summary(rule_findings: list[dict]) -> dict:
    """Strict-POC summary: per-rule pass/fail/warning rollup + pass rate."""
    fail = {r["rule_id"] for r in rule_findings if r["result"] == "fail" and r["rule_id"]}
    warn = {r["rule_id"] for r in rule_findings if r["result"] == "warning" and r["rule_id"]}
    warn -= fail  # a failed rule isn't also counted as a warning
    failed, warned = len(fail), len(warn)
    passed = max(0, 18 - failed - warned)
    return {
        "rules_evaluated": 18,
        "rules_failed": failed,
        "rules_warned": warned,
        "rules_passed": passed,
        "pass_rate_pct": round(passed / 18 * 100, 1),
        "overall": "fail" if failed else ("warn" if warned else "pass"),
        "total_changes": len([r for r in rule_findings if r["applied"]]),
    }


def _apply_to_paragraph(paragraph, corrections: list[dict], transforms: list[tuple] = ()) -> bool:
    """Apply text replacements to a python-docx paragraph in place. Returns True
    if the paragraph text changed. Rewrites the whole paragraph into its first
    run (clearing the rest) so replacements survive Word's run-splitting.

    ``transforms`` are deterministic document-wide (regex, repl) fixes applied
    after the verbatim substitutions, so rule-based corrections land even when the
    finding's quoted text doesn't match the document literally."""
    if not paragraph.runs:
        return False
    full = "".join(run.text for run in paragraph.runs)
    new = full
    for c in corrections:
        if c["original"] and c["original"] in new:
            new = new.replace(c["original"], c["replacement"])
    for rx, repl in transforms:
        new = rx.sub(repl, new)
    if new == full:
        return False
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def _apply_tracked_changes_to_paragraph(paragraph, corrections: list[dict], revision_id_counter: list[int]) -> bool:
    """Apply text corrections to a python-docx paragraph as native tracked changes (revisions).
    Returns True if any tracked changes were applied."""
    if not paragraph.runs:
        return False

    from docx.oxml.shared import OxmlElement, qn

    full_text = "".join(run.text for run in paragraph.runs)

    # Find all occurrences of any correction's "original" text
    matches = []
    for c in corrections:
        orig = c["original"]
        if not orig:
            continue
        start = 0
        while True:
            pos = full_text.find(orig, start)
            if pos == -1:
                break
            matches.append((pos, pos + len(orig), c))
            start = pos + len(orig)

    if not matches:
        return False

    # Filter out overlapping matches
    matches.sort(key=lambda x: x[0])
    filtered_matches = []
    last_end = 0
    for start, end, c in matches:
        if start >= last_end:
            filtered_matches.append((start, end, c))
            last_end = end

    if not filtered_matches:
        return False

    p_el = paragraph._element
    pPr = p_el.pPr

    p_el.clear()
    if pPr is not None:
        p_el.append(pPr)

    last_idx = 0
    author = "Klara AI"
    date = "2026-06-05T12:00:00Z"

    for start, end, c in filtered_matches:
        # 1. Text before the match
        if start > last_idx:
            before_text = full_text[last_idx:start]
            r = OxmlElement('w:r')
            t = OxmlElement('w:t')
            t.text = before_text
            if before_text.startswith(' ') or before_text.endswith(' '):
                t.set(qn('xml:space'), 'preserve')
            r.append(t)
            p_el.append(r)

        # 2. Deletion of the original text
        del_el = OxmlElement('w:del')
        del_el.set(qn('w:id'), str(revision_id_counter[0]))
        del_el.set(qn('w:author'), author)
        del_el.set(qn('w:date'), date)
        revision_id_counter[0] += 1

        r_del = OxmlElement('w:r')
        del_text = OxmlElement('w:delText')
        del_text.text = c["original"]
        if c["original"].startswith(' ') or c["original"].endswith(' '):
            del_text.set(qn('xml:space'), 'preserve')
        r_del.append(del_text)
        del_el.append(r_del)
        p_el.append(del_el)

        # 3. Insertion of the replacement text
        ins_el = OxmlElement('w:ins')
        ins_el.set(qn('w:id'), str(revision_id_counter[0]))
        ins_el.set(qn('w:author'), author)
        ins_el.set(qn('w:date'), date)
        revision_id_counter[0] += 1

        r_ins = OxmlElement('w:r')
        t_ins = OxmlElement('w:t')
        t_ins.text = c["replacement"]
        if c["replacement"].startswith(' ') or c["replacement"].endswith(' '):
            t_ins.set(qn('xml:space'), 'preserve')
        r_ins.append(t_ins)
        ins_el.append(r_ins)
        p_el.append(ins_el)

        last_idx = end

    # 4. Text after the last match
    if last_idx < len(full_text):
        after_text = full_text[last_idx:]
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = after_text
        if after_text.startswith(' ') or after_text.endswith(' '):
            t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p_el.append(r)

    return True


def generate_processed_docx(file_path: str, corrections: list[dict]) -> bytes:
    """Open the original .docx, apply the corrections, return the corrected
    document as bytes. If there are no corrections the original is returned
    unchanged (so the download still succeeds)."""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    if corrections:
        transforms = _rule_transforms_for(corrections)
        for p in doc.paragraphs:
            _apply_to_paragraph(p, corrections, transforms)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _apply_to_paragraph(p, corrections, transforms)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _corrections_signature(corrections: list[dict]) -> str:
    """Stable fingerprint of a correction set, so the cached AI-corrected file is
    rewritten only when the corrections actually change."""
    import hashlib
    import json

    # The leading version tag invalidates cached corrected files whenever the
    # correction-application logic changes (not just the correction data).
    payload = json.dumps(
        ["v2", [[c.get("original"), c.get("replacement"), c.get("rule_id")] for c in corrections]],
        sort_keys=True,
    )
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()[:12]


def build_ai_corrected_file(
    document_id: str, original_path: str, corrections: list[dict]
) -> str:
    """Apply the AI corrections to the pristine original and cache the result as a
    .docx — this is the "AI" side of "Compare original vs AI". Returns the path to
    the corrected file, or ``original_path`` unchanged when there are no
    corrections (so the diff comes out empty rather than erroring).

    The AI's corrections live in the database as findings and are never written to
    the working document on disk, so diffing original-vs-working-file showed
    nothing. Diffing original-vs-this instead surfaces the AI's changes regardless
    of whether a reviewer has manually edited the document yet.

    The file is rewritten only when the correction set changes, so the mtime-keyed
    page-render cache (``build_comparison_images``) stays warm across requests."""
    if not original_path or not os.path.exists(original_path) or not corrections:
        return original_path
    out_dir = os.path.join(os.path.dirname(original_path), ".ai_corrected")
    corrected = os.path.join(out_dir, f"{document_id}.docx")
    marker = corrected + ".sig"
    sig = _corrections_signature(corrections)
    if os.path.exists(corrected) and os.path.exists(marker):
        try:
            with open(marker) as f:
                if f.read().strip() == sig:
                    return corrected
        except Exception:
            pass
    try:
        os.makedirs(out_dir, exist_ok=True)
        data = generate_processed_docx(original_path, corrections)
        with open(corrected, "wb") as f:
            f.write(data)
        with open(marker, "w") as f:
            f.write(sig)
        return corrected
    except Exception as e:
        logger.error(f"AI-corrected docx generation failed: {e}")
        return original_path


def generate_revisioned_docx(file_path: str, corrections: list[dict]) -> bytes:
    """Open the original .docx, apply the corrections as tracked changes (revisions),
    return the revisioned document as bytes."""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    revision_id_counter = [1]
    if corrections:
        for p in doc.paragraphs:
            _apply_tracked_changes_to_paragraph(p, corrections, revision_id_counter)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _apply_tracked_changes_to_paragraph(p, corrections, revision_id_counter)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_comparison(file_path: str, corrections: list[dict]) -> dict:
    """Produce before/after HTML for the side-by-side viewer.

    ``before_html`` is the original preview. ``after_html`` applies each
    correction and wraps the inserted text in ``<span class="ins">``.
    ``redline_html`` keeps the original struck-through next to the replacement.
    Text is HTML-escaped the same way ``docx_to_html`` escapes paragraph text
    so the substring matches line up. (``rule_findings`` and ``summary`` are
    attached by the router from the full finding set.)"""
    before_html = docx_to_html(file_path)
    after_html = before_html
    redline_html = before_html

    for c in corrections:
        orig = html_mod.escape(c["original"])
        repl = html_mod.escape(c["replacement"])
        if orig and orig in after_html:
            after_html = after_html.replace(
                orig, f'<span class="ins">{repl}</span>'
            )
        if orig and orig in redline_html:
            redline_html = redline_html.replace(
                orig,
                f'<span class="del">{orig}</span><span class="ins">{repl}</span>',
            )

    return {
        "before_html": before_html,
        "after_html": after_html,
        "redline_html": redline_html,
        "changes": corrections,
    }


# ── Real document-vs-original redline ─────────────────────────────────────────
# Unlike build_comparison (which paints the AI's *suggested* corrections onto the
# original), this diffs the pristine original against the *actual current
# document* — i.e. the reviewer's manual edits plus the AI changes they accepted,
# with pending AI suggestions resolved away (see docx_to_resolved_paragraphs).
# Deletions are wrapped in <span class="del"> (rendered red, struck through) and
# insertions in <span class="ins"> (rendered green) by the viewer's stylesheet.

_WORD_RE = re.compile(r'\S+|\s+')


def _tokenize_words(s: str) -> list[str]:
    """Split into word and whitespace tokens so spacing survives the diff and
    only the genuinely changed words get marked."""
    return _WORD_RE.findall(s or "")


def _wrap(cls: str, text: str) -> str:
    if not text:
        return ""
    return f'<span class="{cls}">{html_mod.escape(text)}</span>'


def _inline_word_diff(original: str, current: str) -> tuple[str, int, int]:
    """Word-level redline of one paragraph. Returns (html, n_ins, n_del) where
    the counts are of changed word-tokens (whitespace-only changes don't count)."""
    o, c = _tokenize_words(original), _tokenize_words(current)
    sm = difflib.SequenceMatcher(a=o, b=c, autojunk=False)
    parts: list[str] = []
    ins = dele = 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            parts.append(html_mod.escape("".join(o[i1:i2])))
        elif tag == 'delete':
            parts.append(_wrap("del", "".join(o[i1:i2])))
            dele += sum(1 for t in o[i1:i2] if t.strip())
        elif tag == 'insert':
            parts.append(_wrap("ins", "".join(c[j1:j2])))
            ins += sum(1 for t in c[j1:j2] if t.strip())
        elif tag == 'replace':
            parts.append(_wrap("del", "".join(o[i1:i2])))
            parts.append(_wrap("ins", "".join(c[j1:j2])))
            dele += sum(1 for t in o[i1:i2] if t.strip())
            ins += sum(1 for t in c[j1:j2] if t.strip())
    return "".join(parts), ins, dele


def _word_count(s: str) -> int:
    return sum(1 for t in _tokenize_words(s) if t.strip())


def _ptag(inner: str, idx: int | None = None, page: int | None = None) -> str:
    """A diff paragraph. Changed paragraphs carry a stable ``data-change-index``
    and ``data-page`` (matched across both columns) so the viewer can jump
    change-to-change and skip pages with no changes."""
    if idx is None:
        return f'<p>{inner or "&nbsp;"}</p>'
    return (
        f'<p data-change-index="{idx}" data-page="{page}" class="chg">'
        f'{inner or "&nbsp;"}</p>'
    )


def build_change_diff(original_path: str, current_path: str) -> dict:
    """Produce a before/after redline of the pristine original vs the current
    document. ``before_html`` is the (change-tagged) original; ``after_html`` and
    ``redline_html`` are the redline (deletions red + struck, insertions green).
    Each changed paragraph is tagged with a change index + estimated page number,
    and ``changes_index`` lists every change in document order so the viewer can
    navigate between them and skip unchanged pages."""
    orig_paras = docx_to_resolved_paragraphs(original_path)
    curr_paras = docx_to_resolved_paragraphs(current_path)
    orig_pages = estimate_page_map(original_path)
    curr_pages = estimate_page_map(current_path)

    def opage(i: int) -> int:
        return orig_pages[i] if 0 <= i < len(orig_pages) else 1

    def cpage(j: int) -> int:
        return curr_pages[j] if 0 <= j < len(curr_pages) else 1

    esc = html_mod.escape
    sm = difflib.SequenceMatcher(
        a=[p.strip() for p in orig_paras],
        b=[p.strip() for p in curr_paras],
        autojunk=False,
    )

    before_out: list[str] = []   # left column — pristine original (change-tagged)
    after_out: list[str] = []    # right column — redline
    changes_index: list[dict] = []
    ins_words = del_words = 0
    k = 0  # running change index

    def record(kind: str, page: int, preview: str):
        nonlocal k
        changes_index.append({
            "index": k, "page": page, "type": kind,
            "preview": (preview or "").strip()[:90],
        })
        k += 1

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            for o in range(i1, i2):
                before_out.append(_ptag(esc(orig_paras[o])))
                after_out.append(_ptag(esc(orig_paras[o])))
        elif tag == 'delete':
            for o in range(i1, i2):
                pg = opage(o)
                before_out.append(_ptag(esc(orig_paras[o]), k, pg))
                after_out.append(_ptag(_wrap("del", orig_paras[o]), k, pg))
                del_words += _word_count(orig_paras[o])
                record("del", pg, orig_paras[o])
        elif tag == 'insert':
            for c in range(j1, j2):
                pg = cpage(c)
                before_out.append(_ptag("", k, pg))  # left placeholder keeps sync
                after_out.append(_ptag(_wrap("ins", curr_paras[c]), k, pg))
                ins_words += _word_count(curr_paras[c])
                record("ins", pg, curr_paras[c])
        elif tag == 'replace':
            span = max(i2 - i1, j2 - j1)
            for d in range(span):
                oi, cj = i1 + d, j1 + d
                if oi < i2 and cj < j2:
                    inner, ni, nd = _inline_word_diff(orig_paras[oi], curr_paras[cj])
                    if ni or nd:
                        pg = cpage(cj)
                        before_out.append(_ptag(esc(orig_paras[oi]), k, pg))
                        after_out.append(_ptag(inner, k, pg))
                        ins_words += ni
                        del_words += nd
                        record("mod", pg, curr_paras[cj])
                    else:
                        before_out.append(_ptag(esc(orig_paras[oi])))
                        after_out.append(_ptag(inner))
                elif oi < i2:
                    pg = opage(oi)
                    before_out.append(_ptag(esc(orig_paras[oi]), k, pg))
                    after_out.append(_ptag(_wrap("del", orig_paras[oi]), k, pg))
                    del_words += _word_count(orig_paras[oi])
                    record("del", pg, orig_paras[oi])
                else:
                    pg = cpage(cj)
                    before_out.append(_ptag("", k, pg))
                    after_out.append(_ptag(_wrap("ins", curr_paras[cj]), k, pg))
                    ins_words += _word_count(curr_paras[cj])
                    record("ins", pg, curr_paras[cj])

    before_html = "".join(before_out)
    redline_html = "".join(after_out)
    pages_total = max(
        max(curr_pages, default=1), max(orig_pages, default=1), 1
    )
    changed_pages = sorted({c["page"] for c in changes_index})
    return {
        "before_html": before_html,
        "after_html": redline_html,
        "redline_html": redline_html,
        "changes_index": changes_index,
        "changed_pages": changed_pages,
        "pages_total": pages_total,
        "change_stats": {
            "insertions": ins_words,
            "deletions": del_words,
            "paragraphs_changed": len(changes_index),
            "has_changes": bool(changes_index),
            "pages_total": pages_total,
            "changed_pages": changed_pages,
        },
    }


# ── Rendered redline page IMAGES ──────────────────────────────────────────────
# The HTML diff above loses the document's real layout (pleading paper, line
# numbers, fonts). For a faithful "Compare original vs AI", we bake the diff into
# a real .docx as colored/shaded runs (insertions = green + underline + green
# shading; deletions = red + strikethrough + red shading), render it to page
# PNGs with LibreOffice, and detect which pages actually carry a highlight. The
# original is rendered too, for a true side-by-side. Results are cached on disk
# keyed by the original+current mtimes.

COMPARISON_DPI = 150
_RL_GREEN_FILL, _RL_PINK_FILL = "C6EFCE", "FFC7CE"   # light green / light pink
_RL_GREEN_FG, _RL_PINK_FG = "006100", "C00000"       # dark green / dark red


def _seg_diff(orig: str, curr: str) -> list[tuple[str, str]]:
    """Word-level segments of one paragraph: (kind, text), kind ∈ equal/del/ins."""
    o, c = _tokenize_words(orig), _tokenize_words(curr)
    sm = difflib.SequenceMatcher(a=o, b=c, autojunk=False)
    segs: list[tuple[str, str]] = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            segs.append(("equal", "".join(o[i1:i2])))
        elif tag == "delete":
            segs.append(("del", "".join(o[i1:i2])))
        elif tag == "insert":
            segs.append(("ins", "".join(c[j1:j2])))
        elif tag == "replace":
            segs.append(("del", "".join(o[i1:i2])))
            segs.append(("ins", "".join(c[j1:j2])))
    return segs


def _rl_run(qn, OxmlElement, deepcopy, base, text: str, kind: str):
    """Build a <w:r> run carrying ``text`` with redline formatting for ``kind``,
    inheriting font/size/bold/italic from ``base`` rPr (canonical child order)."""
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if base is not None:
        for tag in ("w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs"):
            e = base.find(qn(tag))
            if e is not None:
                rpr.append(deepcopy(e))
    if kind == "del":
        rpr.append(OxmlElement("w:strike"))
    if kind in ("del", "ins"):
        col = OxmlElement("w:color")
        col.set(qn("w:val"), _RL_PINK_FG if kind == "del" else _RL_GREEN_FG)
        rpr.append(col)
    if base is not None:
        for tag in ("w:sz", "w:szCs"):
            e = base.find(qn(tag))
            if e is not None:
                rpr.append(deepcopy(e))
    if kind == "ins":
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    if kind in ("del", "ins"):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), _RL_PINK_FILL if kind == "del" else _RL_GREEN_FILL)
        rpr.append(shd)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _rewrite_para_redline(qn, OxmlElement, deepcopy, p_el, segs):
    """Replace a paragraph's runs with redline-formatted segment runs, keeping
    its pPr (alignment, spacing, numbering) and the first run's base formatting."""
    first_r = p_el.find(qn("w:r"))
    base = first_r.find(qn("w:rPr")) if first_r is not None else None
    base = deepcopy(base) if base is not None else None
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)
    for kind, text in segs:
        if text:
            p_el.append(_rl_run(qn, OxmlElement, deepcopy, base, text, kind))


def generate_redline_docx(original_path: str, current_path: str) -> bytes:
    """Open the pristine original and bake the original→current diff into it as
    colored/shaded redline runs, preserving the document's real formatting.
    Returns the redline .docx bytes."""
    from docx import Document as DocxDocument
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from copy import deepcopy

    orig_paras = docx_to_resolved_paragraphs(original_path)
    curr_paras = docx_to_resolved_paragraphs(current_path)
    doc = DocxDocument(original_path)
    p_elements = doc.element.body.findall(".//" + qn("w:p"))

    def new_ins_para(ref, text):
        new_p = OxmlElement("w:p")
        pPr = ref.find(qn("w:pPr")) if ref is not None else None
        if pPr is not None:
            new_p.append(deepcopy(pPr))
        new_p.append(_rl_run(qn, OxmlElement, deepcopy, None, text, "ins"))
        return new_p

    sm = difflib.SequenceMatcher(
        a=[p.strip() for p in orig_paras],
        b=[p.strip() for p in curr_paras],
        autojunk=False,
    )
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag == "replace":
            span = max(i2 - i1, j2 - j1)
            for d in range(span):
                oi, cj = i1 + d, j1 + d
                if oi < i2 and cj < j2 and oi < len(p_elements):
                    segs = _seg_diff(orig_paras[oi], curr_paras[cj])
                    if any(k != "equal" for k, _ in segs):
                        _rewrite_para_redline(qn, OxmlElement, deepcopy, p_elements[oi], segs)
                elif oi < i2 and oi < len(p_elements):
                    _rewrite_para_redline(qn, OxmlElement, deepcopy, p_elements[oi], [("del", orig_paras[oi])])
                elif cj < j2 and p_elements:
                    ref = p_elements[min(i1, len(p_elements) - 1)]
                    ref.addprevious(new_ins_para(ref, curr_paras[cj]))
        elif tag == "delete":
            for oi in range(i1, i2):
                if oi < len(p_elements):
                    _rewrite_para_redline(qn, OxmlElement, deepcopy, p_elements[oi], [("del", orig_paras[oi])])
        elif tag == "insert" and p_elements:
            before = i1 < len(p_elements)
            ref = p_elements[i1] if before else p_elements[-1]
            for cj in range(j1, j2):
                np = new_ins_para(ref, curr_paras[cj])
                ref.addprevious(np) if before else ref.addnext(np)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _detect_highlights(png_bytes: bytes) -> tuple[int, int]:
    """Count (green-insertion, pink-deletion) highlight pixels on a rendered page.
    Tuned to the exact shading fills and resistant to the document's own colored
    text (e.g. orange template notes) via tight bounds on the blue channel."""
    try:
        import io
        from PIL import Image

        im = Image.open(io.BytesIO(png_bytes)).convert("RGB")
        px = im.load()
        w, h = im.size
        g = p = 0
        for y in range(0, h, 3):
            for x in range(0, w, 3):
                R, G, B = px[x, y]
                if G >= 210 and 150 <= R <= 212 and 160 <= B <= 212 and G - R >= 18 and G - B >= 18:
                    g += 1
                elif R >= 235 and 158 <= G <= 205 and 165 <= B <= 215 and R - G >= 35 and abs(B - G) <= 22:
                    p += 1
        return g, p
    except Exception as e:
        logger.warning(f"highlight detection failed: {e}")
        return 0, 0


def _cmp_cache_key(document_id: str, original_path: str, current_path: str) -> str:
    om = int(os.path.getmtime(original_path)) if original_path and os.path.exists(original_path) else 0
    cm = int(os.path.getmtime(current_path)) if current_path and os.path.exists(current_path) else 0
    return f"{document_id}_{om}_{cm}"


def _cmp_cache_dir(current_path: str, key: str) -> str:
    return os.path.join(os.path.dirname(current_path), ".cmp_cache", key)


def build_comparison_images(document_id: str, original_path: str, current_path: str) -> dict:
    """Render the redline (and original) to page PNGs, detect which pages carry a
    change, and cache everything on disk. Idempotent: a cached manifest short-
    circuits re-rendering until either file changes (mtime-keyed)."""
    import glob
    import json

    from src.utils.doc_render import render_available, render_docx_to_images

    empty = {"images_available": False, "pages_total": 0, "pages": [],
             "changed_pages": [], "has_original": False, "version": ""}
    if not current_path or not os.path.exists(current_path) or not render_available():
        return empty

    key = _cmp_cache_key(document_id, original_path, current_path)
    cache_dir = _cmp_cache_dir(current_path, key)
    manifest = os.path.join(cache_dir, "manifest.json")
    if os.path.exists(manifest):
        try:
            with open(manifest) as f:
                return json.load(f)
        except Exception:
            pass

    base_cache = os.path.join(os.path.dirname(current_path), ".cmp_cache")
    for d in glob.glob(os.path.join(base_cache, f"{document_id}_*")):
        if os.path.abspath(d) != os.path.abspath(cache_dir):
            shutil.rmtree(d, ignore_errors=True)
    os.makedirs(cache_dir, exist_ok=True)

    try:
        redline_bytes = generate_redline_docx(original_path, current_path)
    except Exception as e:
        logger.error(f"redline docx generation failed: {e}")
        return empty
    redline_docx = os.path.join(cache_dir, "redline.docx")
    with open(redline_docx, "wb") as f:
        f.write(redline_bytes)

    redline_imgs = render_docx_to_images(redline_docx, max_pages=50, dpi=COMPARISON_DPI)
    if not redline_imgs:
        logger.warning("redline render produced no pages")
        return empty
    orig_imgs = []
    if original_path and os.path.exists(original_path):
        orig_imgs = render_docx_to_images(original_path, max_pages=50, dpi=COMPARISON_DPI)

    pages = []
    for i, data in enumerate(redline_imgs, 1):
        with open(os.path.join(cache_dir, f"redline-{i}.png"), "wb") as f:
            f.write(data)
        g, p = _detect_highlights(data)
        pages.append({"page": i, "has_changes": bool(g > 40 or p > 40),
                      "insertions": g, "deletions": p})
    for i, data in enumerate(orig_imgs, 1):
        with open(os.path.join(cache_dir, f"original-{i}.png"), "wb") as f:
            f.write(data)

    result = {
        "images_available": True,
        "pages_total": len(redline_imgs),
        "pages": pages,
        "changed_pages": [pg["page"] for pg in pages if pg["has_changes"]],
        "has_original": bool(orig_imgs),
        # Cache-busting token: changes whenever the original or current file
        # changes, so the browser never serves a stale page image.
        "version": key,
    }
    try:
        with open(manifest, "w") as f:
            json.dump(result, f)
    except Exception:
        pass
    return result


def comparison_page_file(
    document_id: str, original_path: str, current_path: str, page: int, side: str
) -> str | None:
    """Path to a cached comparison page PNG (rendering on demand if needed)."""
    if not current_path or not os.path.exists(current_path):
        return None
    key = _cmp_cache_key(document_id, original_path, current_path)
    cache_dir = _cmp_cache_dir(current_path, key)
    name = ("redline" if side == "redline" else "original") + f"-{page}.png"
    fp = os.path.join(cache_dir, name)
    if not os.path.exists(fp):
        build_comparison_images(document_id, original_path, current_path)
    return fp if os.path.exists(fp) else None
